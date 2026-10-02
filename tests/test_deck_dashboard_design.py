#!/usr/bin/env python3
"""tests/test_deck_dashboard_design.py — the investor deck, held to the dashboard.

The deck's palette was repainted to Slate & Cyan while its GRAMMAR stayed where
it was: light boxes on a dark page, figures in a proportional face, a cyan rail
on every slide, tiles that threw their movement away, bands in balance order,
and a pound sign whatever the book was denominated in. A funder reading the pack
beside the dashboard saw two products.

Where it can, this file reads the dashboard's OWN source — ``index.css`` for the
tokens, ``stratOrder.ts`` and ``FundedSnapshotPanel.tsx`` for the rules — so the
two surfaces cannot drift apart again without a test noticing.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import pytest

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from mi_agent_pptx.pptx_theme import THEME                 # noqa: E402

_UI = _ROOT / "frontend" / "mi-agent-ui" / "src"


def _css_token(name: str) -> str:
    css = (_UI / "index.css").read_text(encoding="utf-8")
    m = re.search(rf"{re.escape(name)}:\s*(#[0-9a-fA-F]{{6}})", css)
    assert m, f"{name} is not defined in index.css"
    return m.group(1).lower()


# --------------------------------------------------------------------------- #
# Tokens: read from index.css, not restated.
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("field,token", [
    ("bg_page", "--surface-dashboard"),
    ("bg_cover", "--color-app-ground"),
    ("bg_panel_alt", "--color-navy-800"),
    ("bg_inset", "--color-navy-850"),
    ("bg_well", "--color-navy-950"),
    ("line", "--color-line"),
    ("line_soft", "--color-line-soft"),
    ("line_strong", "--color-line-strong"),
    ("ink_100", "--color-ink-100"),
    ("ink_200", "--color-ink-200"),
    ("ink_300", "--color-ink-300"),
    ("ink_400", "--color-ink-400"),
    ("ink_500", "--color-ink-500"),
    ("ink_600", "--color-ink-600"),
    ("mint", "--color-mint-400"),
    ("rose", "--color-rose-400"),
    ("cyan_500", "--color-cyan-500"),
    ("bar_neutral", "--color-navy-500"),
    ("peri", "--color-cyan-400"),
])
def test_every_deck_token_is_the_dashboards(field, token):
    assert getattr(THEME, field).lower() == _css_token(token), (field, token)


def test_cards_sit_darker_than_the_slide_and_tiles_one_step_above_cards():
    """The dashboard's elevation grammar: dark wells set into the slate surface,
    raised tiles a step lighter than the well. The deck used to invert it."""
    def lum(hx):
        r, g, b = (int(hx[i:i + 2], 16) for i in (1, 3, 5))
        return 0.2126 * r + 0.7152 * g + 0.0722 * b
    assert lum(THEME.bg_panel) < lum(THEME.bg_page)
    assert lum(THEME.bg_panel) < lum(THEME.bg_panel_alt) < lum(THEME.bg_page)
    assert lum(THEME.bg_well) < lum(THEME.bg_panel)


def test_a_neutral_bar_stays_visible_on_a_card():
    """A waterfall's opening bar was drawn in navy-800 — a SURFACE colour. Once
    cards became the dark well it all but disappeared, leaving its label
    floating over nothing. A neutral quantity needs real contrast with the card."""
    def lum(hx):
        r, g, b = (int(hx[i:i + 2], 16) for i in (1, 3, 5))
        return 0.2126 * r + 0.7152 * g + 0.0722 * b
    assert lum(THEME.bar_neutral) - lum(THEME.bg_panel) > 35
    src = (_ROOT / "mi_agent_pptx" / "chart_resolver.py").read_text(encoding="utf-8")
    assert '"base": theme.navy' not in src and '"base": self.theme.navy' not in src


def test_figures_are_set_in_a_monospace_face():
    """``.t-figure`` is monospace with tabular figures."""
    css = (_UI / "index.css").read_text(encoding="utf-8")
    figure_rule = css[css.index(".t-figure"):css.index("}", css.index(".t-figure"))]
    assert "font-mono" in figure_rule
    assert THEME.font_figure and THEME.font_figure != THEME.font_sans


# --------------------------------------------------------------------------- #
# Colour is state, not decoration.
# --------------------------------------------------------------------------- #

def _deck_builder(tmp_path, **data_fields):
    from mi_agent_pptx.deck import DeckBuilder, DeckContext
    from mi_agent_pptx.mi_api import DashboardData

    data = DashboardData(client_id="qa", run_id="mi_2026_06")
    data.reporting_date = "2026-06-30"
    for k, v in data_fields.items():
        setattr(data, k, v)
    ctx = DeckContext(client_name="QA", as_of_date="2026-06-30",
                      run_dir=str(tmp_path), work_dir=str(tmp_path / "c"))
    return DeckBuilder(data, ctx)


def _fills_and_text_colours(slide):
    fills, text = [], []
    for shp in slide.shapes:
        try:
            if shp.fill.type is not None:
                fills.append(str(shp.fill.fore_color.rgb).lower())
        except Exception:
            pass
        if shp.has_text_frame:
            for para in shp.text_frame.paragraphs:
                for run in para.runs:
                    try:
                        text.append(str(run.font.color.rgb).lower())
                    except Exception:
                        pass
    return fills, text


def test_a_slide_header_carries_no_accent_colour(tmp_path):
    """Every slide used to carry a full-height cyan rail and a cyan strapline —
    the one colour the dashboard reserves for selection, on every page."""
    b = _deck_builder(tmp_path)
    s = b._slide()
    b._header(s, "Title", "Subtitle", accent=THEME.peri)
    fills, text = _fills_and_text_colours(s)
    cyan = THEME.peri.lstrip("#").lower()
    assert cyan not in fills, "an accent rail is still drawn"
    assert cyan not in text, "the strapline is still cyan"


def test_no_text_in_the_deck_is_set_in_the_accent(tmp_path):
    """Read across the handlers' source: cyan may colour a data series, never a
    label, a title or a figure."""
    src = (_ROOT / "mi_agent_pptx" / "deck.py").read_text(encoding="utf-8")
    assert "color=self.theme.peri" not in src


# --------------------------------------------------------------------------- #
# The KPI tile is the dashboard's StatTile.
# --------------------------------------------------------------------------- #

def _tile_slide(tmp_path, tile):
    from pptx.util import Inches
    b = _deck_builder(tmp_path)
    s = b._slide()
    b._tile(s, Inches(1), Inches(1), Inches(3), Inches(1.34), tile)
    return s


def _all_text(slide):
    return "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)


def test_a_tile_shows_its_movement_and_context(tmp_path):
    s = _tile_slide(tmp_path, {"label": "Current funded balance", "value": "£104.8MM",
                               "delta": "+£3.7MM", "deltaIntent": "positive",
                               "hint": "+3.7% vs prior run"})
    text = _all_text(s)
    assert "+£3.7MM" in text and "+3.7% vs prior run" in text


@pytest.mark.parametrize("intent,token", [("positive", "mint"), ("negative", "rose")])
def test_the_rail_takes_a_colour_only_where_there_is_a_direction(tmp_path, intent, token):
    s = _tile_slide(tmp_path, {"label": "x", "value": "1", "delta": "+1",
                               "deltaIntent": intent})
    fills, _ = _fills_and_text_colours(s)
    assert getattr(THEME, token).lstrip("#").lower() in fills


def test_a_tile_without_movement_has_the_neutral_rail(tmp_path):
    s = _tile_slide(tmp_path, {"label": "x", "value": "50.4%"})
    fills, _ = _fills_and_text_colours(s)
    assert THEME.line_strong.lstrip("#").lower() in fills
    assert THEME.mint.lstrip("#").lower() not in fills


def test_the_figure_is_in_the_figure_face(tmp_path):
    s = _tile_slide(tmp_path, {"label": "Loans funded", "value": "318"})
    faces = [run.font.name for sh in s.shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs for run in p.runs if run.text == "318"]
    assert faces == [THEME.font_figure]


def test_ordinary_figures_land_on_one_size():
    """"£104.8MM" used to come out smaller than "318" beside it."""
    from mi_agent_pptx.deck import DeckBuilder
    b = object.__new__(DeckBuilder)
    sizes = {b._figure_size(v, 2.5) for v in ("£104.8MM", "318", "50.4%", "£329K")}
    assert len(sizes) == 1, sizes


def test_a_long_text_value_steps_down_rather_than_clipping():
    from mi_agent_pptx.deck import DeckBuilder
    b = object.__new__(DeckBuilder)
    assert b._figure_size("Yorkshire and The Humber", 2.5) < b._figure_size("318", 2.5)


def test_redundant_monthly_change_tiles_follow_the_dashboards_rule():
    """Ported from FundedSnapshotPanel.tsx — checked against that file."""
    tsx = (_UI / "components" / "FundedSnapshotPanel.tsx").read_text(encoding="utf-8")
    assert 'k.id === "mom_balance") return !byId.get("balance")?.delta' in tsx
    assert 'k.id === "mom_loans") return !byId.get("loans")?.delta' in tsx

    from mi_agent_pptx.deck import headline_kpis
    with_deltas = [{"id": "balance", "delta": "+1"}, {"id": "loans", "delta": "+5"},
                   {"id": "mom_loans"}, {"id": "mom_balance"}, {"id": "nneg_risk"}]
    assert [k["id"] for k in headline_kpis(with_deltas)] == ["balance", "loans", "nneg_risk"]
    without = [{"id": "balance"}, {"id": "loans"}, {"id": "mom_loans"}]
    assert [k["id"] for k in headline_kpis(without)] == ["balance", "loans", "mom_loans"]


# --------------------------------------------------------------------------- #
# Bars read in the dashboard's order.
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("given,expected", [
    (["20-30%", "40-50%", "50-60%", "70-80%", "80-90%", "60-70%", "30-40%"],
     ["20-30%", "30-40%", "40-50%", "50-60%", "60-70%", "70-80%", "80-90%"]),
    (["70-75", "85+", "60-65", "55-60", "Unknown"], ["55-60", "60-65", "70-75", "85+", "Unknown"]),
    (["North West", "South East", "East of England", "London"],
     ["East of England", "London", "North West", "South East"]),
    (["2019.0", "2022.0", "2020.0"], ["2019", "2020", "2022"]),
    (["<20%", "20-30%", ">=100%"], ["<20%", "20-30%", ">=100%"]),
])
def test_bars_read_in_the_dashboards_order(given, expected):
    """The governed display order both surfaces consume
    (``mi_agent_api.presentation``; its parity with the dashboard is held in
    ``tests/test_presentation_parity.py``)."""
    from mi_agent_api.presentation import order_bars
    assert [b["label"] for b in order_bars([{"label": g} for g in given])] == expected


def test_ordering_never_alters_a_value():
    from mi_agent_api.presentation import order_bars
    out = order_bars([{"label": "40-50%", "balance": 2.0}, {"label": "20-30%", "balance": 9.0}])
    assert {(b["label"], b["balance"]) for b in out} == {("40-50%", 2.0), ("20-30%", 9.0)}


def test_the_deck_has_no_second_ordering_owner():
    """The deck once carried its own port of stratOrder.ts beside the shared
    module, and the publication gate then disagreed with the page it drew."""
    assert not (_ROOT / "mi_agent_pptx" / "strat_order.py").exists()


# --------------------------------------------------------------------------- #
# Currency is the book's.
# --------------------------------------------------------------------------- #

@pytest.fixture
def euro():
    from mi_agent_api import currency
    token = currency._CURRENCY_CODE.set("EUR")
    yield
    currency._CURRENCY_CODE.reset(token)


def test_compact_figures_use_the_governed_currency(euro):
    from mi_agent_pptx.metric_resolver import compact_currency
    assert compact_currency(104_800_000) == "€104.8MM"


def test_prose_uses_the_governed_currency(euro):
    """Shared with the dashboard's observations — fixed at the source."""
    from mi_agent_api.insight_generators import money, signed_money
    assert money(2_700_000) == "€2.7m"
    assert signed_money(-807_000) == "−€807k"


@pytest.mark.parametrize("module", ["concentration", "movement", "watchlist"])
def test_no_private_pound_formatter_survives(module):
    src = (_ROOT / "mi_agent_pptx" / f"{module}.py").read_text(encoding="utf-8")
    assert 'f"£' not in src, f"{module}.py still hard-codes a pound sign"


def test_a_sterling_book_reads_exactly_as_before():
    from mi_agent_api.insight_generators import money
    from mi_agent_pptx.metric_resolver import compact_currency
    assert compact_currency(104_800_000) == "£104.8MM"
    assert money(2_700_000) == "£2.7m"


def test_the_deck_puts_the_books_currency_in_force():
    src = (_ROOT / "mi_agent_pptx" / "mi_api.py").read_text(encoding="utf-8")
    assert "data.currency_code = _currency.resolve_currency_code(" in src
    assert "_currency.use_currency(data.currency_code)" in src
    cli = (_ROOT / "mi_agent_pptx" / "cli.py").read_text(encoding="utf-8")
    assert "_currency.use_currency(data.currency_code)" in cli


# --------------------------------------------------------------------------- #
# Covenant formatting.
# --------------------------------------------------------------------------- #

def test_a_limit_carries_its_governed_operator():
    from mi_agent_pptx import concentration as C
    assert C.format_limit({"limit": 30, "unit": "percent", "operator": "max"}) == "≤ 30.00%"
    assert C.format_limit({"limit": 5, "unit": "percent", "operator": "min"}) == "≥ 5.00%"


def test_headroom_carries_its_unit():
    from mi_agent_pptx import concentration as C
    assert C.format_headroom(15.79, "percent", dp=2) == "15.79pp"
    assert C.format_headroom(None, "percent") == "—"


def test_a_value_beside_its_limit_does_not_round_onto_it():
    from mi_agent_pptx import concentration as C
    assert C.format_measure(29.96, "percent", dp=2) == "29.96%"
    assert C.format_measure(29.96, "percent") == "30.0%", "1dp default unchanged"


# --------------------------------------------------------------------------- #
# Charts.
# --------------------------------------------------------------------------- #

def test_a_rate_series_is_drawn_from_zero():
    from mi_agent_pptx.render import zero_anchored_limits
    lo, hi = zero_anchored_limits([0.5025, 0.5043, 0.5060])
    assert lo == 0.0 and hi > 0.506


def test_a_series_that_can_go_negative_keeps_its_own_axis():
    from mi_agent_pptx.render import zero_anchored_limits
    assert zero_anchored_limits([-2.0, 1.5]) is None


# --------------------------------------------------------------------------- #
# The prior period is the dashboard's.
# --------------------------------------------------------------------------- #

def test_the_deck_finds_the_prior_run_the_way_the_snapshot_route_does(tmp_path):
    """On a book delivered as runs the deck found NO prior period, so every
    movement the dashboard prints under a KPI was blank in the pack."""
    # Through the ``tests`` package, never by putting tests/ on sys.path: that
    # left tests/operations_control shadowing the real package for every test
    # collected after this one.
    from tests import test_deck_generation_route as T
    from mi_agent_pptx.mi_api import _prior_from_runs

    root = T._write_runs(tmp_path / "runs")
    df, run_id, rd = _prior_from_runs(str(root), T.CLIENT, "mi_2026_06")
    assert df is not None and not df.empty
    assert run_id == "mi_2026_05"


def test_a_blob_root_keeps_the_dated_cut_path():
    from mi_agent_pptx.mi_api import _prior_from_runs
    assert _prior_from_runs("blob://processed-v2/out", "c", "mi_2026_06") == (None, None, None)


def test_the_movement_bridge_opens_where_the_snapshot_compares():
    """The bridge was always meant to open at the snapshot's prior period; with
    no prior it fell back to the earliest period and printed four months of
    movement beside one-month headline figures."""
    src = (_ROOT / "mi_agent_pptx" / "mi_api.py").read_text(encoding="utf-8")
    assert "out_root=out_root, run_id=rid)" in src
